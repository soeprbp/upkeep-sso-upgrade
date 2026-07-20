from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "UpKeep_SSO_Ready_To_Send_Emails.docx"


EMAILS = [
    {
        "title": "Employee Awareness Notice",
        "to": "All employees who may hear about the UpKeep change",
        "subject": "Upcoming UpKeep sign-in update",
        "body": [
            "Hi team,",
            "We are preparing an update to how UpKeep users sign in. UpKeep authentication will be moving toward Microsoft Entra ID single sign-on, which means users will eventually sign in with their Welch work account instead of a separate UpKeep password.",
            "This change is being handled in phases. IT will configure and test the sign-in connection first, then a small pilot group will validate the experience before the broader rollout.",
            "There is no action needed from most employees right now. Users who are part of the pilot or rollout waves will receive separate instructions before their sign-in process changes.",
            "The goal is to make access easier, improve security, and reduce password-related support issues.",
            "Thank you.",
        ],
    },
    {
        "title": "UpKeep User Announcement",
        "to": "UpKeep users",
        "subject": "UpKeep is moving to Microsoft Entra ID sign-in",
        "body": [
            "Hi team,",
            "We are upgrading UpKeep sign-in to Microsoft Entra ID single sign-on (SSO). Going forward, you will use your Welch work account to access UpKeep instead of maintaining a separate UpKeep password.",
            "The change will happen in phases over the next two weeks. A small pilot group will test the new sign-in flow first, then the remaining users will move in waves before final cutover.",
            "What this means for you:",
            "- Use Continue with SSO and your Welch email address when signing in to UpKeep.",
            "- If you use the mobile app, you may be asked for the company ID provided by IT.",
            "- If you are in a later rollout wave, you will receive a follow-up message with your timing and any extra steps.",
            "Native password login will remain available during testing and the pilot period. Once the rollout is complete, UpKeep will use SSO only.",
            "If you have trouble signing in or your UpKeep account uses a non-Welch email address, please contact the helpdesk.",
            "Thank you for your patience while we make this transition.",
        ],
    },
    {
        "title": "Pilot User Instructions",
        "to": "UpKeep SSO pilot users",
        "subject": "UpKeep pilot sign-in instructions",
        "body": [
            "Hi pilot team,",
            "You are in the first group testing the new UpKeep SSO sign-in flow.",
            "When you are ready, please follow these steps:",
            "- Open UpKeep and select Continue with SSO on the main login screen.",
            "- Enter your Welch email address when prompted.",
            "- Complete the Microsoft sign-in steps if you are redirected to Entra ID.",
            "- If you are using the mobile app, enter the company ID if the app asks for it.",
            "- After sign-in, confirm that you can reach your normal UpKeep workspace and complete your usual work order tasks.",
            "If anything fails, send the exact error message and a screenshot to IT so we can correct it quickly. Native login will stay available during the pilot, so use the backup path if you cannot complete SSO.",
            "Please reply after your first successful login and include any friction you noticed, even if you worked around it.",
            "Thank you for helping us validate the rollout before it goes wider.",
        ],
    },
    {
        "title": "Management Update",
        "to": "Leadership and upper management",
        "subject": "UpKeep SSO upgrade underway to improve security and simplify access",
        "body": [
            "Hi team,",
            "We are beginning an UpKeep authentication upgrade that moves users from local UpKeep passwords to Microsoft Entra ID single sign-on.",
            "The goal is to make access easier for users, improve security by using company-managed credentials, and reduce the long-term support burden tied to password resets and account maintenance.",
            "We are handling this as a phased rollout rather than a single cutover. IT will complete the Entra setup, a pilot group will validate the sign-in experience, and the remaining users will move in waves after the process is confirmed.",
            "During the rollout, users will receive instructions before their sign-in process changes. Native login will remain available until the pilot and wave testing are stable.",
            "The key success factors are accurate account matching, a clear help path for users, and a controlled final cutover once the pilot confirms the flow is working as expected.",
            "We will share progress updates as the rollout moves through pilot, waves, and final cutover.",
            "Thank you.",
        ],
    },
]


def set_run_font(run, size=11, color="000000", bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_para(doc, text="", size=11, color="000000", bold=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run)


def add_email(doc, email, index):
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run(f"{index}. {email['title']}")

    for label, value in (("To", email["to"]), ("Subject", email["subject"])):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        label_run = p.add_run(f"{label}: ")
        set_run_font(label_run, bold=True, color="1F4D78")
        value_run = p.add_run(value)
        set_run_font(value_run)

    add_para(doc, "", after=2)

    for line in email["body"]:
        if line.startswith("- "):
            add_bullet(doc, line[2:])
        else:
            add_para(doc, line)


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("UpKeep SSO Ready-To-Send Emails")
    set_run_font(run, size=20, color="1F4D78", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Employee, user, pilot, and management communications prepared from the rollout plan.")
    set_run_font(run, size=11, color="555555")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("UpKeep SSO Communications")
    set_run_font(footer_run, size=9, color="666666")

    for index, email in enumerate(EMAILS, start=1):
        add_email(doc, email, index)

    return doc


def main():
    doc = build_doc()
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()

